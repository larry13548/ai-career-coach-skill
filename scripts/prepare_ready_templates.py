#!/usr/bin/env python3
"""Development-only: normalize the nine supplied styles for deterministic filling."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[1]
DESIGNS={
'T01':('252525',None,'FFFFFF','right'),
'T02':('187FAA','EFF3F7','187FAA','left'),
'T03':('E59470','39485B','FFFFFF','left'),
'T04':('006EBB','F4F8FC','006EBB','left'),
'T05':('252525',None,'FFFFFF','right'),
'T06':('2D4154','EEF2F5','2D4154','left'),
'T07':('466999',None,'FFFFFF','left'),
'T08':('43556E',None,'FFFFFF','right'),
'T09':('60ADB5','34434F','FFFFFF','left')}
def shade(element,fill):
 pr=element.get_or_add_tcPr() if element.tag==qn('w:tc') else element.get_or_add_pPr()
 s=OxmlElement('w:shd');s.set(qn('w:fill'),fill);pr.append(s)
def font(style,size,color,bold=False):
 style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string(color);style.font.bold=bold
 f=style.element.get_or_add_rPr().get_or_add_rFonts();f.set(qn('w:eastAsia'),'Microsoft YaHei')
def marker(container,text,style='Resume Text'):
 p=container.add_paragraph(text,style);return p

def prepare(t,design):
 accent,side,side_color,photo_side=design
 doc=Document(ROOT/'assets/catalog'/f'{t}.docx')
 body=doc.element.body
 for c in list(body):body.remove(c)
 # Discard all old visual/content parts; retained reference stays untouched.
 for rid,rel in list(doc.part.rels.items()):
  if any(rel.reltype.endswith('/'+x) for x in ['image','header','footer','comments','footnotes','endnotes','customXml']):doc.part.drop_rel(rid)
 sect=OxmlElement('w:sectPr');body.append(sect)
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.35);sec.bottom_margin=Cm(1.35);sec.left_margin=Cm(1.35);sec.right_margin=Cm(1.35)
 sec.header_distance=Cm(.5);sec.footer_distance=Cm(.5)
 for name,size,col,bold in [('Resume Text',11,'333333',False),('Resume Heading',12,accent,True),('Resume Entry',11,accent,True),('Resume Name',23,accent,True),('Resume Side',10.5,side_color,False),('Resume Side Heading',12,side_color,True),('Resume Side Name',19,side_color,True)]:
  style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
  font(style,size,col,bold);pf=style.paragraph_format;pf.line_spacing=1.25;pf.space_after=Pt(5);pf.space_before=Pt(0);pf.widow_control=True;pf.keep_with_next=False
  if 'Heading' in name:pf.space_before=Pt(12);pf.space_after=Pt(6);pf.keep_with_next=True
  if name=='Resume Entry':pf.space_before=Pt(5);pf.keep_with_next=True
 font(doc.styles['Normal'],11,'333333');doc.styles['Normal'].paragraph_format.space_after=Pt(0)
 if t=='T02':
  borders=OxmlElement('w:pgBorders');borders.set(qn('w:offsetFrom'),'page')
  for edge in ['top','left','bottom','right']:
   el=OxmlElement('w:'+edge)
   for key,value in [('val','single'),('sz','12'),('space','12'),('color',accent)]:el.set(qn('w:'+key),value)
   borders.append(el)
  sec._sectPr.append(borders)
 if t=='T04':
  marker(doc,'RESUME  / / /','Resume Heading')
 if t=='T05':
  h=marker(doc,'个人简历   / / / / / /','Resume Heading');shade(h._p,accent);h.runs[0].font.color.rgb=RGBColor(255,255,255)
 if t=='T08':
  h=marker(doc,'RESUME','Resume Heading');h.paragraph_format.space_before=Pt(0)
 if t in ['T03','T09']:
  font(doc.styles['Resume Heading'],12,'39485B',True);font(doc.styles['Resume Entry'],11,'39485B',True)
 if side:
  table=doc.add_table(rows=1,cols=2);table.autofit=False
  for col,width in zip(table.columns,[5.4,12.9]):col.width=Cm(width)
  left,right=table.rows[0].cells
  for cell,width in [(left,5.4),(right,12.9)]:
   cell.width=Cm(width)
   pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'170');e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run('').font.size=Pt(1)
  shade(left._tc,side)
  marker(left,'{{PHOTO}}','Resume Side').alignment=WD_ALIGN_PARAGRAPH.CENTER
  marker(left,'{{NAME}}','Resume Side Name')
  marker(left,'{{TARGET}}','Resume Side')
  marker(left,'{{CONTACT}}','Resume Side')
  marker(left,'{{SUMMARY}}','Resume Side')
  if t=='T03':
   h=marker(right,'PERSONAL RESUME','Resume Heading');shade(h._p,accent);h.runs[0].font.color.rgb=RGBColor(255,255,255)
  marker(right,'{{SECTIONS}}')
 else:
  table=doc.add_table(rows=1,cols=2);table.autofit=False
  widths=[3.5,14.8] if photo_side=='left' else [14.8,3.5]
  for col,cell,width in zip(table.columns,table.rows[0].cells,widths):col.width=Cm(width);cell.width=Cm(width)
  info=table.cell(0,1 if photo_side=='left' else 0);photo=table.cell(0,0 if photo_side=='left' else 1)
  info.paragraphs[0].text='{{NAME}}';info.paragraphs[0].style=doc.styles['Resume Name']
  marker(info,'{{TARGET}}');marker(info,'{{CONTACT}}')
  if t=='T07':
   shade(info.paragraphs[0]._p,accent)
   for run in info.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255)
  photo.paragraphs[0].text='{{PHOTO}}';photo.paragraphs[0].style=doc.styles['Resume Text']
  marker(doc,'{{SUMMARY}}');marker(doc,'{{SECTIONS}}')
  if t in ['T01','T05','T07','T08']:
   st=doc.styles['Resume Heading'];font(st,12,'FFFFFF',True)
   pr=st.element.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),accent);pr.append(sh)
   st.paragraph_format.left_indent=Cm(.15)
 # Identity of normalized selected template travels with its source asset.
 doc.core_properties.title=t+' 自动填充版';doc.core_properties.author='';doc.core_properties.subject='Source-derived flow layout 1.5'
 out=ROOT/'assets/ready';out.mkdir(exist_ok=True);doc.save(out/f'{t}.docx')
if __name__=='__main__':
 for t,d in DESIGNS.items():prepare(t,d)
