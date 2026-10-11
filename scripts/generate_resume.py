#!/usr/bin/env python3
"""Fill selected prepared template once; optional one-engine render within a shared deadline."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
from docx import Document
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
LIMIT=300
TARGET=180
ALLOWED={'education','work','internship','project','campus','skills','awards'}

def start(path):
    state={'started_at':time.time(),'limit_seconds':LIMIT,'target_seconds':TARGET,'attempts':0}
    with Path(path).open('x',encoding='utf-8') as f:json.dump(state,f)
    return state

def remaining(state):
    elapsed=time.time()-state['started_at']
    if elapsed<0:raise ValueError('clock moved backwards; cannot verify elapsed time')
    return max(0,LIMIT-elapsed)

def text(value,label):
    if not isinstance(value,str) or not value.strip():raise ValueError(label+' must be nonempty text')
    if any(ord(c)<32 and c not in '\t\n\r' for c in value):raise ValueError('invalid XML character')
    return value

def validate(data):
    if not isinstance(data,dict) or data.get('confirmed') is not True:raise ValueError('confirm full content first')
    if data.get('template_id') not in {f'T{i:02d}' for i in range(1,10)}:raise ValueError('select T01–T09 first')
    for k in ['summary_original','summary_additions']:
        if not isinstance(data.get(k),list):raise ValueError(k+' must be a list; preserve original points')
        for v in data[k]:text(v,k)
    if type(data.get('omit_summary',False)) is not bool:raise ValueError('omit_summary must be boolean')
    if data.get('omit_summary') and data.get('summary_omission_confirmed') is not True:raise ValueError('summary omission needs explicit user choice')
    if not isinstance(data.get('personal'),dict):raise ValueError('personal must be an object')
    for k,v in data['personal'].items():
        text(k,'personal key')
        if v is not None and v!='':text(v,'personal value')
    for k in ('name','target','photo'):
        if data.get(k):text(data[k],k)
    if not isinstance(data.get('sections'),list):raise ValueError('sections must be a list')
    for section in data['sections']:
        text(section.get('title'),'section title')
        if section.get('kind') not in ALLOWED and section.get('independent_section_confirmed') is not True:
            raise ValueError('standalone practice/custom section requires explicit confirmation')
        entries=section.get('entries')
        if not isinstance(entries,list):raise ValueError('entries must be a list')
        for ent in entries:
            if not isinstance(ent,dict):raise ValueError('entry must be an object')
            for k in ('heading','period','detail'):
                if ent.get(k):text(ent[k],k)
            if not isinstance(ent.get('bullets',[]),list):raise ValueError('bullets must be a list')
            for v in ent.get('bullets',[]):text(v,'bullet')
            if not any(ent.get(k) for k in ['heading','period','detail','bullets']):raise ValueError('empty entry')

def insert(before,value,style):
    e=OxmlElement('w:p');before._p.addprevious(e);p=Paragraph(e,before._parent)
    p.style=style;p.add_run(value);return p

def fill(data,output,root=ROOT):
    validate(data);output=Path(output)
    if output.exists():raise ValueError('output exists; use a new filename')
    source=root/'assets/ready'/f"{data['template_id']}.docx"
    doc=Document(source)
    if doc.core_properties.title!=data['template_id']+' 自动填充版':raise ValueError('template identity mismatch')
    markers={}
    for e in doc.element.body.iter(qn('w:p')):
        p=Paragraph(e,doc);v=p.text.strip()
        if v.startswith('{{') and v.endswith('}}'):
            if v in markers:raise ValueError('duplicate template slot')
            markers[v]=p
    if set(markers)!={'{{NAME}}','{{TARGET}}','{{CONTACT}}','{{PHOTO}}','{{SUMMARY}}','{{SECTIONS}}'}:raise ValueError('prepared template slots missing')
    namep=markers['{{NAME}}']
    value=data.get('name') or '姓名：待补充'
    if namep.runs:
        namep.runs[0].text=value
        for run in namep.runs[1:]:run.text=''
    else:namep.add_run(value)
    markers['{{TARGET}}'].text='求职意向：'+(data.get('target') or '待补充')
    info=dict(data['personal'])
    for key in ['性别','年龄','手机','邮箱']:info.setdefault(key,'待补充')
    p=markers['{{CONTACT}}'];style=p.style
    lines=[key+'：'+(value or '待补充') for key,value in info.items()]
    if style.name!='Resume Side':
        # Compact identity block without making long contacts overflow.
        merged=[];pending=''
        for line in lines:
            if pending and len(pending)+len(line)<=40:
                merged.append(pending+'   |   '+line);pending=''
            else:
                if pending:merged.append(pending)
                pending=line
        if pending:merged.append(pending)
        lines=merged
    for line in lines:insert(p,line,style)
    p._p.getparent().remove(p._p)
    p=markers['{{PHOTO}}'];p.text=''
    if data.get('photo'):
        from docx.image.image import Image
        img=Image.from_file(data['photo']);maxw=2.7;maxh=3.5
        width=min(maxw,maxh*img.width/img.height)
        p.add_run().add_picture(data['photo'],width=Cm(width))
    else:p.add_run('照片待上传')
    summary=[]
    for line in data['summary_original']+data['summary_additions']:
        if line not in summary:summary.append(line)
    p=markers['{{SUMMARY}}'];side=p.style.name=='Resume Side'
    if not data.get('omit_summary'):
        insert(p,'个人介绍','Resume Side Heading' if side else 'Resume Heading')
        for line in summary or ['待补充']:insert(p,line,'Resume Side' if side else 'Resume Text')
    p._p.getparent().remove(p._p)
    p=markers['{{SECTIONS}}']
    for section in data['sections']:
        if not section['entries']:continue
        insert(p,section['title'],'Resume Heading')
        for entry in section['entries']:
            heading='  |  '.join(entry[k] for k in ['period','heading'] if entry.get(k))
            if heading:insert(p,heading,'Resume Entry')
            if entry.get('detail'):insert(p,entry['detail'],'Resume Text')
            for line in entry.get('bullets',[]):insert(p,'• '+line,'Resume Text')
    p._p.getparent().remove(p._p)
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=(data.get('name') or '待补充')+' 简历'
    buffer=BytesIO();doc.save(buffer)
    # Content check is independent of render capability; no source sample should survive.
    check=Document(BytesIO(buffer.getvalue()));alltext='\n'.join(''.join(t.text or '' for t in e.iter(qn('w:t'))) for e in check.element.body.iter(qn('w:p')))
    expected=[] if data.get('omit_summary') else summary
    for section in data['sections']:
        if section['entries']:expected.append(section['title'])
        for entry in section['entries']:
            expected.extend(entry[k] for k in ['heading','period','detail'] if entry.get(k));expected.extend(entry.get('bullets',[]))
    for value in expected:
        if value not in alltext:raise ValueError('content verification failed')
    if '{{' in alltext:raise ValueError('unfilled slot')
    output.parent.mkdir(parents=True,exist_ok=True)
    with output.open('xb') as f:f.write(buffer.getvalue())
    return {'file':str(output),'template_id':data['template_id'],'content_verified':True,'visual_status':'not_verified'}

def run_command(cmd,state,maximum,cwd=None):
    budget=min(maximum,remaining(state)-3)
    if budget<=0:raise TimeoutError('shared five-minute deadline reached')
    return subprocess.run(cmd,cwd=cwd,timeout=budget,check=True,capture_output=True,text=True)

def render(output,state,soffice=None):
    # One fixed engine only. Never probe Word, QuickLook, Pages, or change installed fonts.
    engine=soffice or shutil.which('soffice') or shutil.which('libreoffice')
    if not engine:return {'visual_status':'not_verified','reason':'LibreOffice unavailable; no alternate engine attempted'}
    qa=Path(str(output)+'.qa');qa.mkdir(exist_ok=False)
    try:
        with tempfile.TemporaryDirectory(prefix='resume-lo-') as tmp:
            run_command([engine,'-env:UserInstallation='+Path(tmp).as_uri(),'--headless','--convert-to','pdf','--outdir',str(qa),str(output)],state,60)
        pdf=qa/(Path(output).stem+'.pdf')
        if not pdf.is_file():raise ValueError('renderer did not produce a PDF')
        images=shutil.which('pdftoppm')
        if not images:return {'visual_status':'awaiting_visual_review','pdf':str(pdf),'reason':'PDF created; PNG converter unavailable'}
        run_command([images,'-scale-to','1400','-png',str(pdf),str(qa/'page')],state,45)
        pages=sorted(str(p) for p in qa.glob('page-*.png'))
        if not pages:raise ValueError('renderer produced no page images')
        return {'visual_status':'awaiting_visual_review','pdf':str(pdf),'pages':pages}
    except (subprocess.TimeoutExpired,subprocess.CalledProcessError,TimeoutError,ValueError,OSError) as exc:
        return {'visual_status':'not_verified','reason':str(exc)[:400]}

def generate(input_path,output,state_path,do_render=False,soffice=None):
    data=json.loads(Path(input_path).read_text());state=json.loads(Path(state_path).read_text())
    if type(state.get('started_at')) not in (float,int) or type(state.get('attempts')) is not int:raise ValueError('invalid session state')
    if remaining(state)<=0:raise TimeoutError('five-minute deadline reached; do not restart session to continue')
    if not 0<=state['attempts']<3:raise ValueError('initial generation plus two corrections already used')
    state['attempts']+=1;Path(state_path).write_text(json.dumps(state))
    if data.get('photo') and not Path(data['photo']).is_absolute():data['photo']=str(Path(input_path).resolve().parent/data['photo'])
    result=fill(data,output)
    if do_render:result.update(render(Path(output).resolve(),state,soffice))
    result['elapsed_seconds']=round(time.time()-state['started_at'],2)
    result['remaining_seconds']=round(remaining(state),2)
    result['target_met']=result['elapsed_seconds']<=TARGET
    result['attempt']=state['attempts']
    # Awaiting visual review is never success; model must inspect returned pages.
    return result

def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='cmd',required=True)
    b=sub.add_parser('start');b.add_argument('state')
    g=sub.add_parser('generate');g.add_argument('input');g.add_argument('output');g.add_argument('--state',required=True);g.add_argument('--render',action='store_true');g.add_argument('--soffice')
    a=p.parse_args()
    try:
        result=start(a.state) if a.cmd=='start' else generate(a.input,a.output,a.state,a.render,a.soffice)
        print(json.dumps(result,ensure_ascii=False,indent=2));return 0
    except (OSError,ValueError,TypeError,KeyError,TimeoutError) as exc:
        print(json.dumps({'error':str(exc)},ensure_ascii=False));return 2
if __name__=='__main__':sys.exit(main())
