import copy
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from resume_docx import inspect,patch

class ResumeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.dir=Path(self.tmp.name);self.src=self.dir/'source.docx';self.out=self.dir/'out.docx'
        doc=Document();doc.add_paragraph('示例姓名');doc.add_paragraph('邮箱：待补充');doc.save(self.src)
        self.edits={'confirmed':True,'template_id':'T07','edits':[{'index':0,'expected':'示例姓名','replacement':'合成候选人'}]}

    def test_patch_preserves_other_parts_and_missing_fields(self):
        patch(self.src,self.edits,self.out)
        self.assertEqual([x['text'] for x in inspect(self.out)],['合成候选人','邮箱：待补充'])
        with ZipFile(self.src) as a,ZipFile(self.out) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))

    def test_confirmation_and_template_required(self):
        for key,value in [('confirmed',False),('confirmed',None),('template_id',None),('template_id','single-column'),('template_id','T10')]:
            data=copy.deepcopy(self.edits);data[key]=value
            with self.assertRaises(ValueError):patch(self.src,data,self.out)
            self.assertFalse(self.out.exists())

    def test_stale_and_duplicate_edits_rejected_without_output(self):
        for data in [dict(self.edits,edits=[dict(self.edits['edits'][0],expected='过期')]),dict(self.edits,edits=self.edits['edits']*2)]:
            with self.assertRaises(ValueError):patch(self.src,data,self.out)
            self.assertFalse(self.out.exists())

    def test_never_overwrites(self):
        before=self.src.read_bytes()
        with self.assertRaises(ValueError):patch(self.src,self.edits,self.src)
        self.assertEqual(before,self.src.read_bytes())
        self.out.write_bytes(b'existing')
        with self.assertRaises(ValueError):patch(self.src,self.edits,self.out)
        self.assertEqual(self.out.read_bytes(),b'existing')

    def test_drawing_paragraph_rejected(self):
        doc=Document(self.src);doc.paragraphs[0].add_run()._r.append(OxmlElement('w:drawing'));doc.save(self.src)
        self.assertTrue(inspect(self.src)[0]['complex'])
        with self.assertRaises(ValueError):patch(self.src,self.edits,self.out)
        self.assertFalse(self.out.exists())

    def test_no_build_entry(self):
        import subprocess
        result=subprocess.run([sys.executable,str(Path(__file__).resolve().parents[1]/'scripts/resume_docx.py'),'build'],capture_output=True)
        self.assertNotEqual(result.returncode,0)
        self.assertIn(b'invalid choice',result.stderr)

if __name__=='__main__':unittest.main()
